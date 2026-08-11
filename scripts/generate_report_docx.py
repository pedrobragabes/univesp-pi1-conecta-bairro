"""Gera a versão editável do relatório final do Conecta Bairro."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "04-relatorio-final.md"
SCREENSHOT = ROOT / "docs" / "imagens" / "tela-inicial.png"
OUTPUT = ROOT / "entregas" / "Relatorio_Final_Conecta_Bairro.docx"

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_FILL = "F4F6F9"
GREEN = "1F6A52"
ORANGE = "EC7040"
MUTED = "5F6B66"
TOTAL_WIDTH = 9360


def set_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def ensure_child(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def apply_table_geometry(table, weights):
    total = sum(weights)
    widths = [round(TOTAL_WIDTH * weight / total) for weight in weights]
    widths[-1] += TOTAL_WIDTH - sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(TOTAL_WIDTH))
    tbl_ind = ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_mar = ensure_child(tc_pr, "w:tcMar")
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                margin = ensure_child(tc_mar, f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    set_font(run, size=9, color=MUTED)


def replace_markers(text):
    text = re.sub(r"\[PREENCHER[^\]]*\]", "a completar após a atividade de campo", text)
    text = text.replace("[nomes]", "integrantes da equipe")
    text = text.replace("[polo]", "polo da equipe")
    return text


def add_inline(paragraph, text):
    text = replace_markers(text)
    token_pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
    for token in token_pattern.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            set_font(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith("*") and token.endswith("*"):
            set_font(paragraph.add_run(token[1:-1]), italic=True)
        elif token.startswith("`") and token.endswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=9, color=DARK_BLUE)
        else:
            set_font(paragraph.add_run(token))


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    callout = doc.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = normal
    callout.font.color.rgb = RGBColor.from_string(INK)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(10)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO"), size=12, bold=True, color=GREEN)
    p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("PJI110 · PROJETO INTEGRADOR EM COMPUTAÇÃO I"), size=10, bold=True, color=MUTED)
    p.paragraph_format.space_after = Pt(72)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("CONECTA BAIRRO"), name="Georgia", size=30, bold=True, color=INK)
    p.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Sistema web para registro e acompanhamento\nde solicitações comunitárias"), name="Georgia", size=18, italic=True, color=GREEN)
    p.paragraph_format.space_after = Pt(72)

    for label, value in (
        ("Integrantes e RAs", "não informados"),
        ("Polo", "não informado"),
        ("Orientador(a)", "não informado"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(f"{label}: "), bold=True)
        set_font(p.add_run(value), color=MUTED)

    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("São Paulo\n2026"), size=11, color=MUTED)
    doc.add_page_break()


def add_front_note(doc):
    p = doc.add_paragraph(style="Callout")
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FDE9DF")
    p_pr.append(shd)
    add_inline(p, "**Versão-base editável.** Os resultados técnicos estão preenchidos com evidências executadas. Identificação, contexto parceiro e validação com usuários permanecem descritos como não realizados; devem ser atualizados pela equipe com dados reais antes da submissão.")


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for idx, value in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(replace_markers(value)), size=9, bold=True, color="FFFFFF")
        set_cell_shading(cell, GREEN)
    set_repeat_table_header(table.rows[0])
    for row_values in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
    weights = {2: [1.6, 2.4], 3: [1.6, 0.8, 1.5], 4: [1.4, 0.9, 1.1, 1.3]}.get(len(rows[0]), [1] * len(rows[0]))
    apply_table_geometry(table, weights)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot(doc):
    if not SCREENSHOT.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(SCREENSHOT), width=Inches(6.2))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", "Tela inicial do Conecta Bairro com painel e acesso ao cadastro")
    caption = doc.add_paragraph("Figura 1 — Tela inicial do protótipo Conecta Bairro.", style="Caption")
    caption.paragraph_format.keep_with_next = True
    source = doc.add_paragraph("Fonte: elaboração da equipe (2026).")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(8)
    for run in source.runs:
        set_font(run, size=9, italic=True, color=MUTED)


def parse_markdown(doc, lines):
    index = 0
    title_count = 0
    screenshot_added = False
    while index < len(lines):
        raw = lines[index].rstrip()
        line = replace_markers(raw)
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            title_count += 1
            if title_count > 1:
                doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            if line.startswith("### 7.4") and not screenshot_added:
                add_screenshot(doc)
                screenshot_added = True
            doc.add_heading(line[4:], level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Callout")
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_FILL)
            p_pr.append(shd)
            add_inline(p, line[2:])
        elif line.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [part.strip() for part in replace_markers(lines[index].strip()).strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    table_rows.append(values)
                index += 1
            if table_rows:
                add_table(doc, table_rows)
            continue
        elif re.match(r"^- ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        else:
            p = doc.add_paragraph()
            add_inline(p, line.rstrip("  "))
        index += 1
    if not screenshot_added:
        add_screenshot(doc)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    configure_styles(doc)
    add_cover(doc)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "CONECTA BAIRRO  ·  PJI110"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_p.runs:
        set_font(run, size=8, bold=True, color=MUTED)
    p_pr = header_p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C8D0CC")
    borders.append(bottom)
    p_pr.append(borders)
    add_page_field(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("RELATÓRIO FINAL"), name="Georgia", size=22, bold=True, color=INK)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Conecta Bairro"), name="Georgia", size=15, italic=True, color=GREEN)
    p.paragraph_format.space_after = Pt(18)
    add_front_note(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    parse_markdown(doc, lines)

    doc.core_properties.title = "Relatório Final — Conecta Bairro"
    doc.core_properties.subject = "Projeto Integrador em Computação I — PJI110"
    doc.core_properties.author = "Equipe do Projeto Integrador — UNIVESP"
    doc.core_properties.keywords = "UNIVESP, PJI110, desenvolvimento web, Express, SQLite"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

